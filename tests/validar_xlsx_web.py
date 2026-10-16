"""Valida los .xlsx que genera la web en el navegador (tests/e2e/excel_web.spec.js).
Uso: python -I tests/validar_xlsx_web.py gastos|libro|asientos <archivo.xlsx>   (sale con 1 si algo falla)"""
import sys
from openpyxl import load_workbook

tipo, ruta = sys.argv[1], sys.argv[2]
wb = load_workbook(ruta)                    # fórmulas
wv = load_workbook(ruta, data_only=True)    # caché
fallos = []


def ok(c, m):
    if not c:
        fallos.append(m)


esperadas = {'gastos': ['Resumen', 'Facturas', 'Por mes', 'Suscripciones y contratos', 'Recomendaciones', 'Libro registro', 'IVA trimestre'],
             'libro': ['Libro registro'], 'asientos': ['Asientos']}[tipo]
ok(wb.sheetnames == esperadas, f'hojas {wb.sheetnames} != {esperadas}')
nform = 0
for ws in wb.worksheets:
    hdr = 11 if ws.title == 'Resumen' else 1
    cab = [c for c in ws[hdr] if c.value is not None]
    ok(len(cab) >= 2, f'{ws.title}: sin cabecera')
    for c in cab:
        ok(str(c.fill.fgColor.rgb).endswith('0F4C5C'), f'{ws.title}!{c.coordinate}: relleno {c.fill.fgColor.rgb}')
        ok(c.font.b and str(c.font.color.rgb).endswith('FFFFFF'), f'{ws.title}!{c.coordinate}: fuente de cabecera')
        ancho = ws.column_dimensions[c.column_letter].width or 0
        txt = 'ene 202' if hasattr(c.value, 'year') or isinstance(c.value, (int, float)) else str(c.value)
        ok(ancho >= min(max(len(txt) + 6, 10), 70) - (0 if ws.title != 'Resumen' else 99), f'{ws.title}!{c.column_letter}: ancho {ancho} < cabecera+6')
    ok(ws.page_setup.orientation == 'landscape' and ws.page_setup.paperSize == 9, f'{ws.title}: impresión A4 horizontal')
    ok(ws.sheet_properties.pageSetUpPr is not None and ws.sheet_properties.pageSetUpPr.fitToPage, f'{ws.title}: fitToPage')
    ok(ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0, f'{ws.title}: 1 página de ancho')
    ok(ws.freeze_panes is not None, f'{ws.title}: sin paneles inmovilizados')
    ok(ws.auto_filter.ref, f'{ws.title}: sin autofiltro')
    for fila in ws.iter_rows():
        for c in fila:
            es_f = isinstance(c.value, str) and c.value.startswith('=')
            if es_f:
                nform += 1
            v = wv[ws.title][c.coordinate].value
            ok(not (isinstance(v, str) and v.startswith('#')), f'{ws.title}!{c.coordinate}: error {v}')
            if es_f:
                ok(v is not None or c.value.startswith('=IF('), f'{ws.title}!{c.coordinate}: fórmula sin caché')


def fmt_col(ws, nombre):
    for c in ws[1]:
        if c.value == nombre:
            return [ws.cell(r, c.column).number_format for r in range(2, ws.max_row)]
    return []


if tipo == 'gastos':
    f = wb['Facturas']
    ult = f.max_row
    ok(ult > 7, f'pocas facturas ({ult - 2})')
    ok(all('€' in x for x in fmt_col(f, 'Total')), 'Total sin formato €')
    ok(all(x == 'dd/mm/yyyy' for x in fmt_col(f, 'Fecha')), 'Fecha sin formato dd/mm/yyyy')
    ok(isinstance(f['A2'].value, (int, float)) or hasattr(f['A2'].value, 'year'), 'fecha no numérica')
    ok(str(f.cell(ult, 10).value).startswith('=SUBTOTAL(109,'), f'fila total Facturas: {f.cell(ult, 10).value}')
    ok(any(isinstance(c.value, str) and 'SUMIFS(Facturas!' in c.value for row in wb['Resumen'].iter_rows() for c in row), 'Resumen sin SUMIFS')
    ok('SUMIFS' in str(wb['Por mes']['B2'].value), 'Por mes sin SUMIFS')
    ok('SUMIFS' in str(wb['IVA trimestre']['E2'].value), 'IVA trimestre sin SUMIFS')
    tot_res = wv['Resumen']['B4'].value
    tot_fac = wv['Facturas'].cell(ult, 10).value
    ok(abs(tot_res - tot_fac) < 0.01, f'Resumen {tot_res} != Facturas {tot_fac}')
    sc = wv['Por mes']
    ok(abs(sc.cell(sc.max_row, sc.max_column).value - tot_fac) < 0.01, 'Por mes total != Facturas')
    li = wv['IVA trimestre']
    ok(abs(li.cell(li.max_row, 7).value - tot_fac) < 0.02, 'IVA trimestre total != Facturas')
    ok(nform > 50, f'pocas fórmulas ({nform})')
else:
    ws = wb.worksheets[0]
    ok(ws.max_row > 2, 'sin filas')
    nombres = [c.value for c in ws[1]]
    money = [i for i, n in enumerate(nombres, 1) if str(n).split()[0].lower() in ('base', 'cuota', 'total', 'debe', 'haber')]
    ok(money, 'sin columnas de importe')
    for i in money:
        ok(all('€' in ws.cell(r, i).number_format for r in range(2, ws.max_row + 1)), f'{nombres[i-1]} sin formato €')
        ok(all(isinstance(ws.cell(r, i).value, (int, float)) or ws.cell(r, i).value is None or str(ws.cell(r, i).value).startswith('=') for r in range(2, ws.max_row + 1)), f'{nombres[i-1]} no numérico')
    fe = [i for i, n in enumerate(nombres, 1) if str(n).lower().startswith('fecha')]
    ok(fe and ws.cell(2, fe[0]).number_format == 'dd/mm/yyyy' and hasattr(ws.cell(2, fe[0]).value, 'year'), 'fecha no es número con formato')
    ok(nform >= 1, 'sin fila de totales con SUM')
print(f'{tipo}: {len(wb.sheetnames)} hojas, {nform} fórmulas, ' + ('OK' if not fallos else f'{len(fallos)} FALLOS'))
for m in fallos[:40]:
    print('  -', m)
sys.exit(1 if fallos else 0)
